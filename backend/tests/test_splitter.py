"""Test slide splitter module."""
import io
import tempfile
import unittest
from pathlib import Path
import pptx
from pptx.util import Inches

from app.ppt.splitter import split_presentation_by_slide


class TestSplitter(unittest.TestCase):
    def test_split_presentation_by_slide(self):
        # Build a 3-slide test presentation
        prs = pptx.Presentation()

        # Slide 1: Title
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "First Master Slide"
        slide1.placeholders[1].text = "Subtitle for Slide 1"

        # Slide 2: Content
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Second Content Slide"
        slide2.placeholders[1].text = "Bullet 1\nBullet 2"

        # Slide 3: Blank with textbox
        slide3 = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tb.text_frame.text = "Third Custom Slide"

        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        with tempfile.TemporaryDirectory() as tmp_dir:
            results = split_presentation_by_slide(
                pptx_source=pptx_bytes,
                presentation_name="Company Deck",
                output_dir=tmp_dir,
            )

            self.assertEqual(len(results), 3)

            # Verify slide 1
            self.assertEqual(results[0].slide_index, 0)
            self.assertEqual(results[0].total_slides, 3)
            self.assertEqual(results[0].title, "First Master Slide")
            self.assertTrue(results[0].pptx_file_path.exists())
            self.assertTrue(results[0].preview_image_path.exists())

            # Check that the resulting presentation has exactly 1 slide
            split_prs_1 = pptx.Presentation(str(results[0].pptx_file_path))
            self.assertEqual(len(split_prs_1.slides), 1)
            self.assertEqual(split_prs_1.slides[0].shapes.title.text, "First Master Slide")

            # Check slide 2
            self.assertEqual(results[1].slide_index, 1)
            self.assertEqual(results[1].title, "Second Content Slide")
            split_prs_2 = pptx.Presentation(str(results[1].pptx_file_path))
            self.assertEqual(len(split_prs_2.slides), 1)
            self.assertEqual(split_prs_2.slides[0].shapes.title.text, "Second Content Slide")

            # Check slide 3
            self.assertEqual(results[2].slide_index, 2)
            self.assertEqual(results[2].title, "Third Custom Slide")
            split_prs_3 = pptx.Presentation(str(results[2].pptx_file_path))
            self.assertEqual(len(split_prs_3.slides), 1)


if __name__ == "__main__":
    unittest.main()

