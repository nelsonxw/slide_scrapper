"""Test detector module."""
import io
import unittest
import pptx
from app.scraper.detector import is_powerpoint_content, is_powerpoint_url_or_header, OLE_MAGIC


class TestDetector(unittest.TestCase):
    def test_detector_valid_pptx(self):
        # Create valid PPTX in memory
        prs = pptx.Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        is_valid, fmt = is_powerpoint_content(pptx_bytes)
        self.assertTrue(is_valid)
        self.assertEqual(fmt, "pptx")

    def test_detector_legacy_ppt(self):
        fake_ppt = OLE_MAGIC + b"\x00" * 1000
        is_valid, fmt = is_powerpoint_content(fake_ppt)
        self.assertTrue(is_valid)
        self.assertEqual(fmt, "ppt")

    def test_detector_invalid_content(self):
        is_valid, fmt = is_powerpoint_content(b"Hello world this is not powerpoint")
        self.assertFalse(is_valid)
        self.assertEqual(fmt, "")

    def test_detector_url_and_header(self):
        self.assertTrue(is_powerpoint_url_or_header("https://example.com/slides.pptx"))
        self.assertTrue(is_powerpoint_url_or_header("https://example.com/download.ppt?id=123"))
        self.assertTrue(
            is_powerpoint_url_or_header(
                "https://example.com/api/download",
                content_type_header="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        )
        self.assertFalse(is_powerpoint_url_or_header("https://example.com/image.png"))


if __name__ == "__main__":
    unittest.main()

