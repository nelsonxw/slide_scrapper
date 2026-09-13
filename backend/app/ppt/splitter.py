"""
PowerPoint Slide Splitter.
Splits a multi-slide presentation into individual, standalone 1-slide .pptx files,
preserving layout, styles, backgrounds, and shapes, and generates preview images.
"""
from __future__ import annotations

import io
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
import pptx

from app.ppt.preview import render_slide_preview


@dataclass
class SplitSlideResult:
    slide_index: int
    total_slides: int
    title: str
    original_presentation_name: str
    pptx_file_path: Path
    preview_image_path: Path
    file_size: int
    slide_filename: str
    preview_filename: str


def split_presentation_by_slide(
    pptx_source: bytes | Path | str,
    presentation_name: str,
    output_dir: Path | str,
) -> list[SplitSlideResult]:
    """
    Splits the presentation into individual 1-slide .pptx files and generates their preview thumbnails.
    """
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    slides_out_dir = output_dir / "slides"
    previews_out_dir = output_dir / "previews"
    slides_out_dir.mkdir(parents=True, exist_ok=True)
    previews_out_dir.mkdir(parents=True, exist_ok=True)

    # Clean presentation name for file naming
    clean_base_name = re.sub(r"[^a-zA-Z0-9_\-]", "_", presentation_name).strip("_") or "presentation"

    # Save original presentation to temp file
    temp_dir = Path(tempfile.mkdtemp(prefix="ppt_split_"))
    try:
        temp_orig_pptx = temp_dir / "source.pptx"
        if isinstance(pptx_source, bytes):
            temp_orig_pptx.write_bytes(pptx_source)
        else:
            shutil.copy2(Path(pptx_source), temp_orig_pptx)

        # Open with python-pptx to get slide count
        try:
            prs = pptx.Presentation(str(temp_orig_pptx))
            total_slides = len(prs.slides)
        except Exception as e:
            # Check if it's a legacy .ppt or COM conversion is required
            if sys.platform == "win32":
                temp_converted = _convert_ppt_to_pptx_com(temp_orig_pptx)
                if temp_converted:
                    prs = pptx.Presentation(str(temp_converted))
                    temp_orig_pptx = temp_converted
                    total_slides = len(prs.slides)
                else:
                    raise ValueError(f"Could not open presentation: {e}")
            else:
                raise ValueError(f"Could not open presentation: {e}")

        if total_slides == 0:
            return []

        results: list[SplitSlideResult] = []

        for idx in range(total_slides):
            slide_title = _get_slide_title(prs.slides[idx], idx)
            slide_file_name = f"{clean_base_name}_slide_{idx + 1}.pptx"
            preview_file_name = f"{clean_base_name}_slide_{idx + 1}.png"

            dest_pptx_path = slides_out_dir / slide_file_name
            dest_preview_path = previews_out_dir / preview_file_name

            # Create standalone 1-slide presentation for slide idx
            _create_single_slide_presentation(temp_orig_pptx, idx, total_slides, dest_pptx_path)

            # Generate slide thumbnail preview
            render_slide_preview(dest_pptx_path, 0, dest_preview_path)

            results.append(
                SplitSlideResult(
                    slide_index=idx,
                    total_slides=total_slides,
                    title=slide_title,
                    original_presentation_name=presentation_name,
                    pptx_file_path=dest_pptx_path,
                    preview_image_path=dest_preview_path,
                    file_size=dest_pptx_path.stat().st_size,
                    slide_filename=slide_file_name,
                    preview_filename=preview_file_name,
                )
            )

        return results

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def _create_single_slide_presentation(
    source_pptx_path: Path,
    keep_index: int,
    total_slides: int,
    dest_path: Path,
) -> None:
    """
    Copies the presentation and removes all slides except `keep_index`.
    """
    # Open fresh copy of presentation
    prs = pptx.Presentation(str(source_pptx_path))
    slide_ids = list(prs.slides._sldIdLst)

    # Remove all slide elements except the one at keep_index (in reverse order to keep indices stable)
    for i in range(len(slide_ids) - 1, -1, -1):
        if i != keep_index:
            rId = slide_ids[i].rId
            try:
                prs.part.drop_rel(rId)
            except Exception:
                pass
            del prs.slides._sldIdLst[i]

    prs.save(str(dest_path))


def _get_slide_title(slide: pptx.slide.Slide, index: int) -> str:
    """
    Extracts the slide title or first prominent text snippet.
    """
    try:
        if slide.shapes.title and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()
    except Exception:
        pass

    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                return text.split("\n")[0][:60]

    return f"Slide {index + 1}"


def _convert_ppt_to_pptx_com(ppt_path: Path) -> Path | None:
    """
    Converts legacy .ppt to .pptx using PowerPoint COM on Windows.
    """
    try:
        import win32com.client
        import pythoncom

        pythoncom.CoInitialize()
        try:
            ppt_app = win32com.client.Dispatch("PowerPoint.Application")
            pres = ppt_app.Presentations.Open(str(ppt_path), ReadOnly=True, Untitled=False, WithWindow=False)
            output_pptx = ppt_path.with_suffix(".pptx")
            # 24 = ppSaveAsOpenXMLPresentation (.pptx format)
            pres.SaveAs(str(output_pptx), 24)
            pres.Close()
            return output_pptx
        finally:
            pythoncom.CoUninitialize()
    except Exception as e:
        print(f"COM conversion failed: {e}")
        return None
